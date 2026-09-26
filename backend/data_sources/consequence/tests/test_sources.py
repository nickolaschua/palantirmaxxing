from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest

import openpyxl

from backend.data_sources.consequence import sources


def workbook(path: Path) -> Path:
    """Two sheets in the MOH layout: title rows, a 'Date' header, then days as Excel serials or datetimes."""
    wb = openpyxl.Workbook()
    hist = wb.active
    hist.title = 'Historical'
    hist.append(['Daily ED Attendances by Hospital'])
    hist.append([])
    hist.append(['Date', 'AH', 'WH'])
    hist.append([46268, 54, None])  # 2026-09-03; WH blank, not zero
    hist.append([46269, 62, 246])
    week = wb.create_sheet('Sheet1')
    week.append(['Date', 'AH', 'WH'])
    week.append([datetime(2026, 9, 8), 73, 235])
    week.append(['Source : Weekly Hospital Submissions'])
    wb.save(path)
    return path


class Workbook(unittest.TestCase):
    def test_daily_series_merges_sheets_and_skips_blanks(self):
        with tempfile.TemporaryDirectory() as tmp:
            series = sources.daily_series(workbook(Path(tmp) / 'emd.xlsx'))
        self.assertEqual(series['AH'], {date(2026, 9, 3): 54.0, date(2026, 9, 4): 62.0, date(2026, 9, 8): 73.0})
        self.assertEqual(series['WH'], {date(2026, 9, 4): 246.0, date(2026, 9, 8): 235.0})

    def test_empty_workbook_is_an_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            openpyxl.Workbook().save(path)
            with self.assertRaises(ValueError):
                sources.daily_series(path)


class Links(unittest.TestCase):
    def test_xlsx_link_is_found_once_and_percent_encoded(self):
        base = 'https://isomer-user-content.by.gov.sg/3/abc/Bed Occupancy Rate_week37Y2026.xlsx'
        html = f'<a href="{base}">x</a><script>"{base}\\"</script><img src="https://isomer-user-content.by.gov.sg/3/l.png">'
        self.assertEqual(sources.xlsx_link(html),
                         'https://isomer-user-content.by.gov.sg/3/abc/Bed%20Occupancy%20Rate_week37Y2026.xlsx')

    def test_missing_or_ambiguous_link_is_an_error(self):
        with self.assertRaises(ValueError):
            sources.xlsx_link('<html></html>')
        two = 'https://isomer-user-content.by.gov.sg/3/a/x.xlsx https://isomer-user-content.by.gov.sg/3/b/y.xlsx'
        with self.assertRaises(ValueError):
            sources.xlsx_link(two)


class Geocode(unittest.TestCase):
    def test_postal_codes_regain_their_leading_zero(self):
        self.assertEqual(sources.postal6('88256'), '088256')
        self.assertEqual(sources.postal6(' 738907 '), '738907')

    def test_pick_postal_needs_an_exact_postal_match(self):
        results = [dict(POSTAL='169609', X='1', Y='2'),
                   dict(POSTAL='169608', X='27710.9', Y='29154.7', ADDRESS='1 HOSPITAL CRESCENT', BUILDING='SGH')]
        hit = sources.pick_postal(results, '169608')
        self.assertEqual((hit['x'], hit['y'], hit['building']), (27710.9, 29154.7, 'SGH'))
        self.assertIsNone(sources.pick_postal(results, '000000'))


class Geocoding(unittest.TestCase):
    RESULTS = [{'BLK_NO': '1', 'ROAD_NAME': 'MY FIRST SKOOL ROAD', 'POSTAL': 'NIL', 'X': '1', 'Y': '1'},
               {'BLK_NO': '1', 'ROAD_NAME': 'BEDOK SOUTH AVENUE 1', 'POSTAL': '460001', 'X': '39173.8', 'Y': '33678.9',
                'ADDRESS': '1 BEDOK SOUTH AVENUE 1'}]

    def test_pick_requires_block_postal_and_street_match(self):
        hit = sources.pick_geocode(self.RESULTS, '1', 'BEDOK STH AVE 1')
        self.assertEqual((hit['x'], hit['postal']), (39173.8, '460001'))
        self.assertIsNone(sources.pick_geocode(self.RESULTS, '2', 'BEDOK STH AVE 1'))
        self.assertIsNone(sources.pick_geocode(self.RESULTS, '1', 'TAMPINES ST 11'))


if __name__ == '__main__':
    unittest.main()
